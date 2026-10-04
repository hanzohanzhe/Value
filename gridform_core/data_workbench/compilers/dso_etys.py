"""DSO geometry and ETYS capability compilers."""

from __future__ import annotations

import hashlib
import json
import math
from collections import defaultdict
from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import load_workbook
from pyproj import CRS, Transformer
from shapely import make_valid
from shapely.geometry import mapping, shape
from shapely.ops import transform


def _hash(payload: object) -> str:
    encoded = json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _ordered_geometry(geometry: object) -> object:
    geo = shape(geometry)
    if geo.geom_type == "MultiPolygon":
        parts = sorted(geo.geoms, key=lambda item: (item.bounds, item.wkb_hex))
        from shapely.geometry import MultiPolygon

        geo = MultiPolygon(parts)
    return mapping(geo)


def compile_dso_geojson(
    payload: Mapping[str, object],
    *,
    source_crs: str,
    overlap_tolerance: float = 1e-12,
) -> dict[str, object]:
    if payload.get("type") != "FeatureCollection" or not isinstance(payload.get("features"), list):
        raise ValueError("DSO source must be a GeoJSON FeatureCollection")
    source_reference = CRS.from_user_input(source_crs)
    display_reference = CRS.from_epsg(4326)
    transformer = Transformer.from_crs(source_reference, display_reference, always_xy=True)
    seen: set[str] = set()
    compiled: list[tuple[str, dict[str, object], object]] = []
    repairs: list[dict[str, str]] = []
    for raw in payload["features"]:  # type: ignore[index]
        if not isinstance(raw, Mapping) or not isinstance(raw.get("properties"), Mapping):
            raise ValueError("DSO source contains an invalid feature")
        props = dict(raw["properties"])
        zone_id = str(props.get("zone_id") or "").strip()
        if not zone_id or zone_id in seen:
            raise ValueError("DSO source contains a blank or duplicate zone ID")
        seen.add(zone_id)
        if not isinstance(raw.get("geometry"), Mapping):
            raise ValueError(f"DSO polygon {zone_id} has no geometry")
        source_geometry = shape(raw["geometry"])
        if source_geometry.geom_type not in {"Polygon", "MultiPolygon"}:
            raise ValueError(f"DSO zone {zone_id} is not polygonal")
        if not source_geometry.is_valid:
            source_geometry = make_valid(source_geometry)
            if source_geometry.geom_type not in {"Polygon", "MultiPolygon"}:
                polygonal = [
                    item for item in getattr(source_geometry, "geoms", ())
                    if item.geom_type in {"Polygon", "MultiPolygon"}
                ]
                if not polygonal:
                    raise ValueError(f"DSO zone {zone_id} could not be repaired as polygonal geometry")
                from shapely.ops import unary_union

                source_geometry = unary_union(polygonal)
            repairs.append({"zone_id": zone_id, "method": "shapely_make_valid"})
        display_geometry = transform(transformer.transform, source_geometry)
        if display_geometry.is_empty:
            raise ValueError(f"DSO zone {zone_id} became empty during CRS conversion")
        feature = {
            "type": "Feature",
            "id": zone_id,
            "properties": {
                **props,
                "zone_id": zone_id,
                "geometry_feature_id": str(props.get("geometry_feature_id") or zone_id),
            },
            "geometry": _ordered_geometry(mapping(display_geometry)),
        }
        zone = {
            "zone_id": zone_id,
            "display_name": str(props.get("display_name") or zone_id),
            "dso_owner": str(props.get("dso_owner") or ""),
            "nation": str(props.get("nation") or ""),
            "geometry_feature_id": str(props.get("geometry_feature_id") or zone_id),
            "zone_kind": "resource_zone",
        }
        if not zone["dso_owner"] or not zone["nation"]:
            raise ValueError(f"DSO zone {zone_id} lacks owner or nation metadata")
        compiled.append((zone_id, feature, display_geometry))

    compiled.sort(key=lambda row: row[0])
    for index, (left_id, _, left) in enumerate(compiled):
        for right_id, _, right in compiled[index + 1 :]:
            overlap = left.intersection(right).area
            if overlap > overlap_tolerance:
                raise ValueError(
                    f"DSO polygons {left_id} and {right_id} have unexpected positive-area overlap"
                )
    result: dict[str, object] = {
        "schema_version": "value.dso-resource-zones/v1",
        "source_crs": source_reference.to_string(),
        "display_crs": "EPSG:4326",
        "resource_zones": [row[1]["properties"] | {"zone_kind": "resource_zone"} for row in compiled],
        "display_geojson": {
            "type": "FeatureCollection",
            "features": [row[1] for row in compiled],
        },
        "repairs": sorted(repairs, key=lambda row: row["zone_id"]),
    }
    result["scientific_sha256"] = _hash(result)
    return result


def _number(value: object, label: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} is not numeric") from exc
    if not math.isfinite(number) or number < 0:
        raise ValueError(f"{label} cannot be negative or non-finite")
    return number


def read_etys_workbook(path: Path, *, sheet_name: str) -> list[dict[str, object]]:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"ETYS workbook has no sheet {sheet_name}")
        rows = workbook[sheet_name].iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise ValueError("ETYS workbook sheet is empty") from exc
        headers = [str(item or "").strip() for item in raw_headers]
        if any(not item for item in headers) or len(headers) != len(set(headers)):
            raise ValueError("ETYS workbook has blank or duplicate headers")
        return [
            dict(zip(headers, row))
            for row in rows
            if any(item not in (None, "") for item in row)
        ]
    finally:
        workbook.close()


def compile_etys_capabilities(
    rows: Sequence[Mapping[str, object]], *, year: int = 2025
) -> dict[str, object]:
    capabilities: dict[str, list[Mapping[str, object]]] = defaultdict(list)
    ignored: list[dict[str, object]] = []
    inventory: set[str] = set()
    for row in rows:
        boundary_id = str(row.get("Boundary") or "").strip()
        if not boundary_id:
            raise ValueError("ETYS row has no boundary ID")
        inventory.add(boundary_id)
        category = str(row.get("Category") or "").strip()
        row_year = int(row.get("Year") or 0)
        if category.casefold() != "capability" or row_year != year:
            ignored.append(
                {
                    "boundary_id": boundary_id,
                    "year": row_year,
                    "category": category,
                    "reason": "not_2025_capability",
                }
            )
            continue
        _number(row.get("Value MW"), "Capability")
        capabilities[boundary_id].append(row)

    boundaries: list[dict[str, object]] = []
    waivers: list[str] = []
    for boundary_id in sorted(capabilities):
        by_direction: dict[str, float] = {}
        for row in capabilities[boundary_id]:
            direction = str(row.get("Direction") or "forward").strip()
            if direction in by_direction:
                raise ValueError(f"Capability for {boundary_id}/{direction} is duplicated")
            by_direction[direction] = _number(row.get("Value MW"), "Capability")
        directions = list(by_direction)
        forward_direction = directions[0]
        forward = by_direction[forward_direction]
        reverse_directions = [item for item in directions if item != forward_direction]
        if reverse_directions:
            reverse_direction = reverse_directions[0]
            reverse = by_direction[reverse_direction]
            method = "independent_source"
        else:
            reverse_direction = f"reverse_of:{forward_direction}"
            reverse = forward
            method = "symmetric_forward_fallback"
            waivers.append(f"etys.{boundary_id}.symmetric_forward_fallback")
        boundaries.append(
            {
                "boundary_id": boundary_id,
                "year": year,
                "category": "Capability",
                "forward_direction": forward_direction,
                "forward_limit_mw": forward,
                "reverse_direction": reverse_direction,
                "reverse_limit_mw": reverse,
                "reverse_limit_method": method,
            }
        )
    result: dict[str, object] = {
        "schema_version": "value.etys-capabilities/v1",
        "boundaries": boundaries,
        "boundary_inventory": sorted(inventory),
        "ignored_rows": sorted(
            ignored, key=lambda row: (str(row["boundary_id"]), int(row["year"]), str(row["category"]))
        ),
        "requested_waivers": sorted(waivers),
    }
    result["scientific_sha256"] = _hash(result)
    return result


def reconcile_boundary_evidence(
    primary: Mapping[str, Mapping[str, object]],
    secondary: Mapping[str, Mapping[str, object]],
) -> dict[str, object]:
    resolved: dict[str, dict[str, object]] = {}
    conflicts: list[dict[str, object]] = []
    inventory: list[dict[str, object]] = []
    for boundary_id in sorted(set(primary) | set(secondary)):
        first = dict(primary.get(boundary_id, {}))
        second = dict(secondary.get(boundary_id, {}))
        first_date = str(first.get("publication_date") or "")
        second_date = str(second.get("publication_date") or "")
        selected: dict[str, object] = {}
        for field in sorted((set(first) | set(second)) - {"publication_date"}):
            first_present = field in first and first[field] not in (None, "")
            second_present = field in second and second[field] not in (None, "")
            if first_present and second_present and first[field] != second[field]:
                if second_date > first_date:
                    value, selected_source = second[field], "secondary"
                else:
                    value, selected_source = first[field], "primary"
                conflicts.append(
                    {
                        "boundary_id": boundary_id,
                        "field": field,
                        "primary_value": first[field],
                        "secondary_value": second[field],
                        "selected_value": value,
                        "selected_source": selected_source,
                    }
                )
                selected[field] = value
            elif first_present:
                selected[field] = first[field]
            elif second_present:
                selected[field] = second[field]
        resolved[boundary_id] = selected
        inventory.append(
            {
                "boundary_id": boundary_id,
                "primary_present": bool(first),
                "secondary_present": bool(second),
            }
        )
    return {
        "schema_version": "value.etys-evidence-reconciliation/v1",
        "resolved": resolved,
        "conflicts": conflicts,
        "inventory": inventory,
    }
