"""Offline normalizers for the official-source Prompt 98 GB candidate.

The functions in this module deliberately stop at reviewable scientific
artifacts.  They do not download data, install a data pack or sign a network
identity.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from typing import Mapping, Sequence
import csv
from datetime import datetime
import heapq
import hashlib
import io
import math
import zipfile

from openpyxl import load_workbook
from pyproj import Transformer
import shapefile
from shapely.geometry import LineString, MultiLineString, mapping, shape
from shapely.ops import split, transform, unary_union
from shapely.strtree import STRtree

from .compilers.dso_etys import compile_dso_geojson
from .compilers.network_candidate import compile_review_candidate
from ..weather_spatialization import ZonalAvailabilityBundle, ZonalAvailabilityProfile


_NATION_BY_AREA = {
    "East England": "England",
    "East Midlands": "England",
    "London": "England",
    "North Wales, Merseyside and Cheshire": "England and Wales",
    "West Midlands": "England",
    "North East England": "England",
    "North West England": "England",
    "North Scotland": "Scotland",
    "South East England": "England",
    "South Wales": "Wales",
    "South West England": "England",
    "Yorkshire": "England",
    "South and Central Scotland": "Scotland",
    "Southern England": "England",
}


def _slug(value: object) -> str:
    text = str(value or "").strip().casefold()
    result = "-".join("".join(char if char.isalnum() else " " for char in text).split())
    if not result:
        raise ValueError("Official geography has a blank stable identity")
    return result


def normalize_dso_source(
    payload: Mapping[str, object], *, source_crs: str = "EPSG:27700"
) -> dict[str, object]:
    """Map NESO's actual DNO fields onto the stable VALUE DSO contract."""

    if payload.get("type") != "FeatureCollection" or not isinstance(payload.get("features"), list):
        raise ValueError("Official DSO object is not a GeoJSON FeatureCollection")
    features: list[dict[str, object]] = []
    for raw in payload["features"]:  # type: ignore[index]
        if not isinstance(raw, Mapping) or not isinstance(raw.get("properties"), Mapping):
            raise ValueError("Official DSO object contains an invalid feature")
        source = raw["properties"]
        area = str(source.get("Area") or "").strip()
        if area not in _NATION_BY_AREA:
            raise ValueError(f"Official DSO area {area!r} lacks an explicit GB nation mapping")
        zone_id = f"dso-{_slug(area)}"
        features.append(
            {
                "type": "Feature",
                "id": zone_id,
                "properties": {
                    "zone_id": zone_id,
                    "display_name": area,
                    "dso_owner": str(source.get("DNO_Full") or source.get("DNO") or "").strip(),
                    "nation": _NATION_BY_AREA[area],
                    "geometry_feature_id": zone_id,
                    "source_feature_id": str(source.get("ID") or ""),
                    "source_short_code": str(source.get("Name") or ""),
                    "source_dno_code": str(source.get("DNO") or ""),
                },
                "geometry": raw.get("geometry"),
            }
        )
    # NESO's licence-area product contains deliberate embedded overlaps (most
    # visibly London within the surrounding East England licence geometry).
    # Preserve the smallest, most specific licence area and carve it out of the
    # larger polygon.  The exact removed area is retained as review evidence.
    source_geometries = {
        str(item["properties"]["zone_id"]): shape(item["geometry"])  # type: ignore[index]
        for item in features
    }
    retained: list[tuple[str, object]] = []
    overlap_resolution: list[dict[str, object]] = []
    by_id = {str(item["properties"]["zone_id"]): item for item in features}  # type: ignore[index]
    for zone_id in sorted(
        source_geometries,
        key=lambda item: (float(source_geometries[item].area), item),  # type: ignore[attr-defined]
    ):
        original = source_geometries[zone_id]
        geometry = original
        for preserved_id, preserved in retained:
            overlap = geometry.intersection(preserved)  # type: ignore[attr-defined]
            if float(overlap.area) <= 0:  # type: ignore[attr-defined]
                continue
            overlap_resolution.append(
                {
                    "carved_zone_id": zone_id,
                    "preserved_zone_id": preserved_id,
                    "overlap_area_source_crs": float(overlap.area),  # type: ignore[attr-defined]
                    "carved_fraction_of_source_area": float(overlap.area) / float(original.area),  # type: ignore[attr-defined]
                    "method": "preserve_smaller_embedded_licence_area",
                }
            )
            geometry = geometry.difference(preserved)  # type: ignore[attr-defined]
        if geometry.is_empty:  # type: ignore[attr-defined]
            raise ValueError(f"DSO overlap resolution removed all of {zone_id}")
        by_id[zone_id]["geometry"] = mapping(geometry)
        retained.append((zone_id, geometry))
    result = compile_dso_geojson(
        {"type": "FeatureCollection", "features": list(by_id.values())},
        source_crs=source_crs,
    )
    result["overlap_resolution"] = sorted(
        overlap_resolution,
        key=lambda item: (str(item["carved_zone_id"]), str(item["preserved_zone_id"])),
    )
    return result


def _header_index(rows: Sequence[Sequence[object]]) -> int:
    for index, row in enumerate(rows):
        values = {str(item or "").strip() for item in row}
        if {"Boundary", "Scenario", "Category"}.issubset(values):
            return index
    raise ValueError("ETYS workbook does not contain the expected chart-data header")


def read_etys_2025_capabilities(
    path: Path, *, sheet_name: str = "ETYS 2025 Chart Data", year: int = 2025
) -> dict[str, object]:
    """Normalize the public ETYS chart workbook's wide capability table."""

    with Path(path).open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"ETYS workbook has no sheet {sheet_name}")
            values = list(workbook[sheet_name].iter_rows(values_only=True))
        finally:
            workbook.close()
    header_row = _header_index(values)
    headers = [str(item or "").strip() for item in values[header_row]]
    try:
        year_index = next(index for index, item in enumerate(headers) if item == str(year))
    except StopIteration as exc:
        raise ValueError(f"ETYS workbook has no {year} capability column") from exc
    indexes = {name: headers.index(name) for name in ("Boundary", "Scenario", "Category")}
    evidence: dict[str, dict[str, float]] = defaultdict(dict)
    ignored: list[dict[str, str]] = []
    for row in values[header_row + 1 :]:
        boundary_id = str(row[indexes["Boundary"]] or "").strip()
        if not boundary_id:
            continue
        category = str(row[indexes["Category"]] or "").strip()
        scenario = str(row[indexes["Scenario"]] or "").strip()
        if category.casefold() != "capability":
            ignored.append({"boundary_id": boundary_id, "category": category})
            continue
        try:
            capability = float(row[year_index])
        except (TypeError, ValueError) as exc:
            raise ValueError(f"ETYS {boundary_id} has a non-numeric {year} capability") from exc
        if not math.isfinite(capability) or capability < 0:
            raise ValueError(f"ETYS {boundary_id} has an invalid {year} capability")
        if scenario in evidence[boundary_id] and evidence[boundary_id][scenario] != capability:
            raise ValueError(f"ETYS {boundary_id} repeats a conflicting {year} capability")
        evidence[boundary_id][scenario] = capability
    boundaries: list[dict[str, object]] = []
    for boundary_id in sorted(evidence):
        values_by_scenario = evidence[boundary_id]
        unique = set(values_by_scenario.values())
        if len(unique) != 1:
            raise ValueError(
                f"ETYS {boundary_id} has conflicting {year} capability across scenarios"
            )
        capability = unique.pop()
        boundaries.append(
            {
                "boundary_id": boundary_id,
                "year": year,
                "forward_limit_mw": capability,
                "reverse_limit_mw": capability,
                "reverse_limit_method": "assumed_symmetric_from_forward",
                "source_direction": "positive_side_to_negative_side",
                "scenario_values": dict(sorted(values_by_scenario.items())),
            }
        )
    return {
        "schema_version": "value.etys-2025-capability-evidence/v1",
        "boundaries": boundaries,
        "ignored_rows": sorted(ignored, key=lambda item: (item["boundary_id"], item["category"])),
    }


def read_etys_boundary_geometry(
    archive: Path, *, source_crs: str = "EPSG:3857"
) -> dict[str, object]:
    """Read a pinned ETYS shapefile ZIP without extracting it to the workspace."""

    transformer = Transformer.from_crs(source_crs, "EPSG:4326", always_xy=True)
    with zipfile.ZipFile(Path(archive)) as bundle:
        shapefile_name = next(
            (name for name in bundle.namelist() if name.casefold().endswith(".shp")),
            None,
        )
        if shapefile_name is None:
            raise ValueError("ETYS geometry archive contains no shapefile")
        stem = shapefile_name[:-4]
        reader = shapefile.Reader(
            shp=BytesIO(bundle.read(shapefile_name)),
            shx=BytesIO(bundle.read(stem + ".shx")),
            dbf=BytesIO(bundle.read(stem + ".dbf")),
        )
        features: list[dict[str, object]] = []
        for source_shape, record in zip(reader.shapes(), reader.records()):
            values = record.as_dict()
            boundary_id = str(values.get("Boundary_n") or "").strip()
            if not boundary_id:
                raise ValueError("ETYS geometry contains a blank boundary ID")
            geometry = transform(transformer.transform, shape(source_shape.__geo_interface__))
            if geometry.geom_type not in {"LineString", "MultiLineString"} or geometry.is_empty:
                raise ValueError(f"ETYS boundary {boundary_id} has invalid line geometry")
            features.append(
                {
                    "type": "Feature",
                    "id": boundary_id,
                    "properties": {
                        "boundary_id": boundary_id,
                        "source_feature_id": str(values.get("ID") or ""),
                        "source_crs": source_crs,
                    },
                    "geometry": mapping(geometry),
                }
            )
        reader.close()
    features.sort(key=lambda item: str(item["properties"]["boundary_id"]))  # type: ignore[index]
    if len({item["properties"]["boundary_id"] for item in features}) != len(features):  # type: ignore[index]
        raise ValueError("ETYS geometry repeats a boundary ID")
    return {"type": "FeatureCollection", "features": features}


def _line_parts(geometry: object) -> list[LineString]:
    value = shape(geometry)
    if isinstance(value, LineString):
        return [value]
    if isinstance(value, MultiLineString):
        return list(value.geoms)
    raise ValueError("Selected ETYS boundary is not linear")


def _split_piece(piece: object, line: object) -> list[object]:
    try:
        result = split(piece, line)  # type: ignore[arg-type]
    except ValueError:
        return [piece]
    polygons = [item for item in result.geoms if item.geom_type in {"Polygon", "MultiPolygon"}]
    return polygons or [piece]


def _extend_line(line: LineString, margin: float) -> LineString:
    coordinates = list(line.coords)
    if len(coordinates) < 2:
        return line
    first, second = coordinates[0], coordinates[1]
    before_last, last = coordinates[-2], coordinates[-1]

    def extension(origin: tuple[float, float], neighbour: tuple[float, float]) -> tuple[float, float]:
        dx = origin[0] - neighbour[0]
        dy = origin[1] - neighbour[1]
        length = math.hypot(dx, dy)
        if length == 0:
            return origin
        return origin[0] + margin * dx / length, origin[1] + margin * dy / length

    return LineString(
        [extension(first, second), *coordinates[1:-1], extension(last, before_last)]
    )


def _merge_slivers(pieces: list[object], minimum_piece_share: float) -> tuple[list[object], int]:
    total = sum(float(item.area) for item in pieces)  # type: ignore[attr-defined]
    retained = [item for item in pieces if float(item.area) >= total * minimum_piece_share]  # type: ignore[attr-defined]
    slivers = [item for item in pieces if item not in retained]
    if not retained:
        return [unary_union(pieces)], max(0, len(pieces) - 1)
    for sliver in slivers:
        target_index = min(
            range(len(retained)),
            key=lambda index: (
                float(sliver.distance(retained[index])),  # type: ignore[attr-defined]
                -float(retained[index].area),  # type: ignore[attr-defined]
                index,
            ),
        )
        retained[target_index] = unary_union([retained[target_index], sliver])
    return retained, len(slivers)


def _positive_side(centroid: object, line: object) -> bool:
    nearest = line.interpolate(line.project(centroid))  # type: ignore[attr-defined]
    return float(centroid.y) >= float(nearest.y)  # type: ignore[attr-defined]


def build_candidate_zones(
    dso_geojson: Mapping[str, object],
    boundary_geojson: Mapping[str, object],
    *,
    selected_boundary_ids: Sequence[str],
    minimum_piece_share: float = 0.01,
) -> dict[str, object]:
    """Split DSO resource zones by a reviewed ETYS subset and propose cut sides."""

    if not 0 <= minimum_piece_share < 0.5:
        raise ValueError("minimum_piece_share must be in [0, 0.5)")
    boundary_features = {
        str(item.get("properties", {}).get("boundary_id") or ""): item  # type: ignore[union-attr]
        for item in boundary_geojson.get("features", [])  # type: ignore[assignment]
    }
    selected = tuple(dict.fromkeys(str(item) for item in selected_boundary_ids))
    if any(item not in boundary_features for item in selected):
        raise ValueError("Selected ETYS boundary is absent from the pinned geometry")
    selected_lines = {
        boundary_id: unary_union(_line_parts(boundary_features[boundary_id]["geometry"]))
        for boundary_id in selected
    }
    raw_dso_features = list(dso_geojson.get("features", []))  # type: ignore[arg-type]
    dso_geometries = [
        shape(item.get("geometry"))
        for item in raw_dso_features
        if isinstance(item, Mapping)
    ]
    if not dso_geometries:
        raise ValueError("Normalized DSO layer contains no geometry")
    bounds = unary_union(dso_geometries).bounds
    margin = max(bounds[2] - bounds[0], bounds[3] - bounds[1]) * 2.0
    split_lines = {
        boundary_id: [_extend_line(line, margin) for line in _line_parts(boundary_features[boundary_id]["geometry"])]
        for boundary_id in selected
    }
    output_features: list[dict[str, object]] = []
    sliver_count = 0
    for raw in raw_dso_features:
        if not isinstance(raw, Mapping) or not isinstance(raw.get("properties"), Mapping):
            raise ValueError("Normalized DSO layer contains an invalid feature")
        parent_id = str(raw["properties"].get("zone_id") or "")
        parent = shape(raw.get("geometry"))
        pieces: list[object] = [parent]
        for boundary_id in selected:
            for boundary_line in split_lines[boundary_id]:
                next_pieces: list[object] = []
                for piece in pieces:
                    next_pieces.extend(_split_piece(piece, boundary_line))
                pieces = next_pieces
        pieces, merged = _merge_slivers(pieces, minimum_piece_share)
        sliver_count += merged
        ordered = sorted(
            pieces,
            key=lambda item: (-float(item.centroid.y), float(item.centroid.x), item.wkb_hex),  # type: ignore[attr-defined]
        )
        for index, piece in enumerate(ordered, start=1):
            child_id = parent_id if len(ordered) == 1 else f"{parent_id}--{index:02d}"
            properties = dict(raw["properties"])
            properties.update(
                {
                    "zone_id": child_id,
                    "display_name": (
                        str(properties.get("display_name") or parent_id)
                        if len(ordered) == 1
                        else f"{properties.get('display_name') or parent_id} {index}"
                    ),
                    "dso_parent_zone_id": parent_id,
                    "geometry_feature_id": child_id,
                    "parent_area_share": float(piece.area) / float(parent.area),  # type: ignore[attr-defined]
                }
            )
            output_features.append(
                {
                    "type": "Feature",
                    "id": child_id,
                    "properties": properties,
                    "geometry": mapping(piece),
                }
            )
    output_features.sort(key=lambda item: str(item["properties"]["zone_id"]))  # type: ignore[index]
    proposals: list[dict[str, object]] = []
    for boundary_id in selected:
        line = selected_lines[boundary_id]
        positive = sorted(
            str(item["properties"]["zone_id"])  # type: ignore[index]
            for item in output_features
            if _positive_side(shape(item["geometry"]).centroid, line)
        )
        negative = sorted(
            str(item["properties"]["zone_id"])  # type: ignore[index]
            for item in output_features
            if str(item["properties"]["zone_id"]) not in positive  # type: ignore[index]
        )
        if not positive or not negative:
            raise ValueError(f"Selected ETYS boundary {boundary_id} does not separate candidate zones")
        proposals.append(
            {
                "boundary_id": boundary_id,
                "positive_zone_ids": positive,
                "negative_zone_ids": negative,
                "orientation_method": "centroid_north_of_nearest_boundary_point",
            }
        )
    return {
        "schema_version": "value.gb-zonal-geometry-candidate/v1",
        "zone_geojson": {"type": "FeatureCollection", "features": output_features},
        "selected_boundary_ids": list(selected),
        "cut_proposals": proposals,
        "merged_sliver_count": sliver_count,
        "zone_count": len(output_features),
    }


def _postcode(value: object) -> str:
    return "".join(str(value or "").upper().split())


def _zone_index(zone_geojson: Mapping[str, object]) -> tuple[list[str], list[object], STRtree]:
    zone_ids: list[str] = []
    geometries: list[object] = []
    for raw in zone_geojson.get("features", []):  # type: ignore[assignment]
        if not isinstance(raw, Mapping) or not isinstance(raw.get("properties"), Mapping):
            raise ValueError("Demand zone layer contains an invalid feature")
        zone_id = str(raw["properties"].get("zone_id") or "")
        if not zone_id:
            raise ValueError("Demand zone layer contains a blank zone ID")
        if raw.get("geometry") is None:
            continue
        zone_ids.append(zone_id)
        geometries.append(shape(raw["geometry"]))
    if not geometries:
        raise ValueError("Demand zone layer contains no polygon geometry")
    return zone_ids, geometries, STRtree(geometries)


def _query_zone(
    longitude: float,
    latitude: float,
    zone_ids: Sequence[str],
    geometries: Sequence[object],
    tree: STRtree,
) -> tuple[str | None, bool]:
    from shapely.geometry import Point

    point = Point(longitude, latitude)
    candidates = tree.query(point)
    indexes: list[int] = []
    for item in candidates:
        try:
            index = int(item)
        except (TypeError, ValueError):
            index = next(
                (position for position, geometry in enumerate(geometries) if geometry.equals(item)),  # type: ignore[attr-defined]
                -1,
            )
        if index >= 0 and geometries[index].covers(point):  # type: ignore[attr-defined]
            indexes.append(index)
    matches = sorted({zone_ids[index] for index in indexes})
    return (matches[0] if matches else None), len(matches) > 1


def compile_postcode_demand_weights(
    consumption_csv: Path,
    onspd_zip: Path,
    zone_geojson: Mapping[str, object],
) -> dict[str, object]:
    """Compile measured GB demand weights from pinned DESNZ and ONSPD objects.

    The implementation streams ONSPD and never extracts its postcode files.
    Full unmatched identifiers are intentionally not written to the review
    artifact; counts, energy and a deterministic bounded sample are retained.
    """

    consumption: dict[str, float] = {}
    aggregate_rows = 0
    with Path(consumption_csv).open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            if str(row.get("Postcode") or "").strip().casefold() == "all postcodes":
                aggregate_rows += 1
                continue
            postcode = _postcode(row.get("Postcode"))
            if not postcode:
                continue
            try:
                value = float(row.get("Total_cons_kwh") or "") / 1000.0
            except (TypeError, ValueError) as exc:
                raise ValueError(f"DESNZ postcode {postcode} has invalid consumption") from exc
            if not math.isfinite(value) or value < 0:
                raise ValueError(f"DESNZ postcode {postcode} has invalid consumption")
            if postcode in consumption:
                raise ValueError(f"DESNZ consumption repeats postcode {postcode}")
            consumption[postcode] = value
    if not consumption:
        raise ValueError("DESNZ consumption contains no individual postcode rows")

    all_zone_ids = sorted(
        str(item.get("properties", {}).get("zone_id") or "")  # type: ignore[union-attr]
        for item in zone_geojson.get("features", [])  # type: ignore[assignment]
    )
    zone_ids, geometries, tree = _zone_index(zone_geojson)
    energy_by_zone = {zone_id: 0.0 for zone_id in all_zone_ids}
    assigned: set[str] = set()
    non_gb: set[str] = set()
    invalid_coordinate: set[str] = set()
    boundary_assignments = 0
    duplicate_live = 0
    with zipfile.ZipFile(Path(onspd_zip)) as bundle:
        members = sorted(
            name for name in bundle.namelist()
            if "/multi_csv/" in name.replace("\\", "/").casefold()
            and name.casefold().endswith(".csv")
        )
        if not members:
            raise ValueError("ONSPD archive contains no split postcode CSV files")
        for member in members:
            with bundle.open(member) as binary:
                with io.TextIOWrapper(binary, encoding="utf-8-sig", newline="") as text:
                    for row in csv.DictReader(text):
                        postcode = _postcode(row.get("pcds") or row.get("pcd7"))
                        if postcode not in consumption:
                            continue
                        terminated = str(row.get("doterm") or "").strip()
                        if terminated and terminated != "99999999":
                            continue
                        country = str(row.get("ctry25cd") or row.get("ctry") or "").strip().upper()
                        if country.startswith("N") or postcode.startswith("BT"):
                            non_gb.add(postcode)
                            continue
                        if not country.startswith(("E", "W", "S")):
                            non_gb.add(postcode)
                            continue
                        if postcode in assigned:
                            duplicate_live += 1
                            continue
                        try:
                            longitude = float(row.get("long") or "")
                            latitude = float(row.get("lat") or "")
                        except (TypeError, ValueError):
                            invalid_coordinate.add(postcode)
                            continue
                        zone_id, on_boundary = _query_zone(
                            longitude, latitude, zone_ids, geometries, tree
                        )
                        if zone_id is None:
                            invalid_coordinate.add(postcode)
                            continue
                        assigned.add(postcode)
                        energy_by_zone[zone_id] += consumption[postcode]
                        boundary_assignments += int(on_boundary)

    excluded_non_gb_mwh = sum(consumption[item] for item in non_gb)
    invalid_coordinate_mwh = sum(consumption[item] for item in invalid_coordinate)
    unmatched = set(consumption).difference(assigned, non_gb, invalid_coordinate)
    unmatched_mwh = sum(consumption[item] for item in unmatched)
    included = sum(energy_by_zone.values())
    if included <= 0:
        raise ValueError("No DESNZ postcode consumption could be assigned to GB zones")
    weights = {zone_id: energy_by_zone[zone_id] / included for zone_id in sorted(energy_by_zone)}
    if weights:
        final = max(weights, key=lambda item: (weights[item], item))
        weights[final] += 1.0 - sum(weights.values())
    return {
        "schema_version": "value.official-postcode-zone-demand/v1",
        "energy_mwh_by_zone": dict(sorted(energy_by_zone.items())),
        "weights": dict(sorted(weights.items())),
        "method_by_zone": {
            zone_id: "measured" if energy_by_zone[zone_id] > 0 else "static_share_fallback"
            for zone_id in sorted(energy_by_zone)
        },
        "coverage": {
            "aggregate_rows_excluded": aggregate_rows,
            "assigned_postcodes": len(assigned),
            "non_gb_postcodes": len(non_gb),
            "invalid_or_outside_postcodes": len(invalid_coordinate),
            "unmatched_postcodes": len(unmatched),
            "duplicate_live_directory_rows": duplicate_live,
            "boundary_assignments": boundary_assignments,
            "unmatched_sample": heapq.nsmallest(100, unmatched),
        },
        "reconciliation": {
            "source_individual_postcode_mwh": sum(consumption.values()),
            "included_gb_mwh": included,
            "excluded_non_gb_mwh": excluded_non_gb_mwh,
            "invalid_or_outside_mwh": invalid_coordinate_mwh,
            "unmatched_mwh": unmatched_mwh,
            "accounted_mwh": included + excluded_non_gb_mwh + invalid_coordinate_mwh + unmatched_mwh,
        },
    }


def normalize_scheme_c_national_demand(
    source: Path, *, period_hours: float = 0.5
) -> list[dict[str, object]]:
    """Freeze the retained VALUE ND field onto a half-hour MWh clock."""

    rows: list[dict[str, object]] = []
    with Path(source).open("r", encoding="utf-8-sig", newline="") as handle:
        for raw in csv.DictReader(handle):
            try:
                date = datetime.strptime(str(raw.get("SETTLEMENT_DATE") or ""), "%d-%b-%Y")
                settlement_period = int(raw.get("SETTLEMENT_PERIOD") or 0)
                demand_mw = float(raw.get("ND") or "")
            except (TypeError, ValueError) as exc:
                raise ValueError("VALUE national demand has an invalid settlement row") from exc
            if not 1 <= settlement_period <= 50 or not math.isfinite(demand_mw) or demand_mw < 0:
                raise ValueError("VALUE national demand has an invalid settlement row")
            rows.append(
                {
                    "period_id": f"{date:%Y-%m-%d}:{settlement_period:02d}",
                    "national_demand_mwh": demand_mw * period_hours,
                }
            )
    period_ids = [str(item["period_id"]) for item in rows]
    if not rows or len(period_ids) != len(set(period_ids)):
        raise ValueError("VALUE national demand clock is empty or duplicated")
    return rows


def _nearest_zone(
    longitude: float,
    latitude: float,
    zone_geojson: Mapping[str, object],
) -> tuple[str | None, bool]:
    from shapely.geometry import Point

    point = Point(longitude, latitude)
    candidates: list[tuple[float, str, bool]] = []
    for raw in zone_geojson.get("features", []):  # type: ignore[assignment]
        if not isinstance(raw, Mapping) or not isinstance(raw.get("properties"), Mapping):
            continue
        if raw.get("geometry") is None:
            continue
        zone_id = str(raw["properties"].get("zone_id") or "")
        geometry = shape(raw["geometry"])
        candidates.append((float(geometry.distance(point)), zone_id, bool(geometry.covers(point))))
    if not candidates:
        return None, False
    distance, zone_id, covered = min(candidates, key=lambda item: (item[0], item[1]))
    return zone_id, covered


def _generator_technology(asset_id: str) -> tuple[str, str | None]:
    if asset_id.startswith("solar_"):
        return "solar", asset_id.removeprefix("solar_")
    if asset_id.startswith("onshore_"):
        return "onshore", asset_id.removeprefix("onshore_")
    if asset_id.startswith("offshore"):
        return "offshore", asset_id
    aliases = {
        "CCGT": "ccgt",
        "OCGT": "ocgt",
        "bio_and_waste": "bio_and_waste",
        "Hydro_natural_flow": "hydro_natural_flow",
        "Nuclear": "nuclear",
    }
    return aliases.get(asset_id, _slug(asset_id)), None


def normalize_scheme_c_fleet(
    zone_geojson: Mapping[str, object],
    *,
    generators: Mapping[str, Mapping[str, object]] | None = None,
    batteries: Mapping[str, Mapping[str, object]] | None = None,
    locations: Mapping[str, Mapping[str, object]] | None = None,
) -> dict[str, object]:
    """Export the copied VALUE fleet without mutating its configuration."""

    if generators is None or batteries is None or locations is None:
        from gridform_core.builtin.scheme_c_1000twh.runtime_compat import config

        generators = generators if generators is not None else config.generators
        batteries = batteries if batteries is not None else config.batteries
        locations = locations if locations is not None else config.locations
    rows: list[dict[str, object]] = []
    for asset_id, values in generators.items():
        technology, location_key = _generator_technology(asset_id)
        if technology in {"onshore", "offshore"}:
            capacity = float(values.get("capacity_multiplier") or 0.0) * 20.0
            capacity_basis = "scheme_c_20mw_multiplier"
        elif technology == "solar":
            capacity = float(values.get("capacity_multiplier") or 0.0)
            capacity_basis = "direct_mw_multiplier"
        else:
            capacity = float(values.get("capacity_limit") or 0.0)
            capacity_basis = "direct_capacity_limit_mw"
        latitude: float | None = None
        longitude: float | None = None
        region = location_key or "not_declared"
        zone_id = "ENGLAND_FALLBACK"
        method = "unlocated_england_fallback"
        if location_key and location_key in locations:
            latitude = float(locations[location_key]["lat"])
            longitude = float(locations[location_key]["lon"])
            resolved, covered = _nearest_zone(longitude, latitude, zone_geojson)
            if resolved:
                zone_id = resolved
                if technology == "offshore":
                    method = "source_coordinate" if covered else "inferred_nearest_coast_dso"
                else:
                    method = "source_coordinate" if covered else "unlocated_england_fallback"
                    if not covered:
                        zone_id = "ENGLAND_FALLBACK"
        rows.append(
            {
                "asset_id": asset_id,
                "technology": technology,
                "capacity_mw": capacity,
                "status": "operating" if capacity > 0 else "inactive_zero_capacity",
                "asset_class": "generator",
                "zone_id": zone_id,
                "share": 1.0,
                "mapping_method": method,
                "region": region,
                "nation": "GB",
                "latitude": latitude,
                "longitude": longitude,
                "capacity_basis": capacity_basis,
            }
        )
    for asset_id, values in batteries.items():
        capacity = float(values.get("per_pool_limit") or 0.0)
        rows.append(
            {
                "asset_id": asset_id,
                "technology": str(values.get("battery_type") or asset_id),
                "capacity_mw": capacity,
                "energy_capacity_mwh": float(values.get("pool_limit") or 0.0),
                "status": "operating" if capacity > 0 else "inactive_zero_capacity",
                "asset_class": "storage",
                "zone_id": "ENGLAND_FALLBACK",
                "share": 1.0,
                "mapping_method": "unlocated_england_fallback",
                "region": "not_declared",
                "nation": "GB",
                "latitude": None,
                "longitude": None,
                "capacity_basis": "per_pool_limit_mw",
            }
        )
    return {
        "schema_version": "value.scheme-c-normalized-fleet-candidate/v1",
        "rows": sorted(rows, key=lambda item: str(item["asset_id"])),
        "source": "copied VALUE runtime compatibility configuration",
        "mutated_retained_source": False,
    }


def _repd_technology(value: object) -> str | None:
    text = str(value or "").casefold()
    if "offshore" in text and "wind" in text:
        return "offshore"
    if "onshore" in text and "wind" in text:
        return "onshore"
    if "solar" in text:
        return "solar"
    if "battery" in text and "pumped" not in text:
        return "battery"
    if "pumped storage" in text:
        return "pumped_hydro"
    if "hydro" in text:
        return "hydro_natural_flow"
    return None


def normalize_repd_source(source: Path) -> dict[str, object]:
    """Normalize the pinned Q2 July 2025 REPD extract for spatial audit."""

    transformer = Transformer.from_crs("EPSG:27700", "EPSG:4326", always_xy=True)
    rows: list[dict[str, object]] = []
    excluded: dict[str, int] = defaultdict(int)
    with Path(source).open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        for raw in csv.DictReader(handle):
            country = str(raw.get("Country") or "").strip()
            if country.casefold() in {"northern ireland", "ni"}:
                excluded["northern_ireland_outside_gb_scope"] += 1
                continue
            technology = _repd_technology(raw.get("Technology Type"))
            if technology is None:
                excluded["technology_outside_prompt98_scope"] += 1
                continue
            try:
                capacity = float(raw.get("Installed Capacity (MWelec)") or "")
            except (TypeError, ValueError):
                excluded["invalid_capacity"] += 1
                continue
            if not math.isfinite(capacity) or capacity <= 0:
                excluded["invalid_capacity"] += 1
                continue
            longitude: float | None = None
            latitude: float | None = None
            easting = raw.get("X-coordinate")
            northing = raw.get("Y-coordinate")
            try:
                longitude, latitude = transformer.transform(float(easting), float(northing))
            except (TypeError, ValueError):
                excluded["missing_or_invalid_coordinate"] += 1
            status_source = str(raw.get("Development Status (short)") or "").strip()
            status = "operational" if status_source.casefold() == "operational" else status_source.casefold().replace(" ", "_")
            rows.append(
                {
                    "asset_id": f"repd:{str(raw.get('Ref ID') or '').strip()}",
                    "site_name": str(raw.get("Site Name") or "").strip(),
                    "technology": technology,
                    "capacity_mw": capacity,
                    "status": status,
                    "status_source": status_source,
                    "region": str(raw.get("Region") or "").strip(),
                    "country": country,
                    "longitude": longitude,
                    "latitude": latitude,
                    "raw_easting": easting,
                    "raw_northing": northing,
                }
            )
    rows.sort(key=lambda item: str(item["asset_id"]))
    ids = [str(item["asset_id"]) for item in rows]
    if any(item == "repd:" for item in ids) or len(ids) != len(set(ids)):
        raise ValueError("REPD normalized identity is blank or duplicated")
    return {
        "schema_version": "value.repd-spatial-candidate/v1",
        "rows": rows,
        "excluded_by_reason": dict(sorted(excluded.items())),
        "source_crs": "EPSG:27700",
        "display_crs": "EPSG:4326",
    }


def compile_network_candidate_payload(
    zone_geojson: Mapping[str, object],
    cut_proposals: Sequence[Mapping[str, object]],
    capabilities: Mapping[str, Mapping[str, object]],
    *,
    fallback_reference: tuple[float, float] = (-1.8904, 52.4862),
) -> dict[str, object]:
    """Compile a provisional graph/cut payload while preserving the owner gate."""

    decisions = {
        "schema_version": "value.network-review-overrides/v1",
        "decisions": [
            {
                "boundary_id": str(row.get("boundary_id") or ""),
                "decision": "accepted",
                "approved_by": "prompt98_offline_candidate_builder",
                "approval_scope": "candidate_mapping_only_not_owner_signoff",
            }
            for row in cut_proposals
        ],
    }
    review_candidate = compile_review_candidate(zone_geojson, cut_proposals, decisions)
    if review_candidate["blocking_reasons"]:
        raise ValueError("Prompt 98 provisional network contains unresolved cut mappings")
    zone_features = list(zone_geojson.get("features", []))  # type: ignore[arg-type]
    central_zone, _ = _nearest_zone(
        fallback_reference[0], fallback_reference[1], zone_geojson
    )
    if not central_zone:
        raise ValueError("England fallback has no deterministic central transport zone")
    fallback = {
        "type": "Feature",
        "id": "ENGLAND_FALLBACK",
        "properties": {
            "zone_id": "ENGLAND_FALLBACK",
            "display_name": "Unlocated England",
            "dso_owner": "VALUE fallback",
            "nation": "England",
            "dso_parent_zone_id": "ENGLAND_FALLBACK",
            "is_unconstrained_fallback": True,
            "island_reason": "unlocated_aggregate_assets_only",
        },
        "geometry": None,
    }
    corridors = [
        {
            "corridor_id": item["corridor_id"],
            "from_zone_id": item["from_zone_id"],
            "to_zone_id": item["to_zone_id"],
            "positive_direction": item["positive_direction"],
            "purpose": "computational_routing",
        }
        for item in review_candidate["corridors"]
    ]
    corridors.append(
        {
            "corridor_id": f"corridor:ENGLAND_FALLBACK--{central_zone}",
            "from_zone_id": "ENGLAND_FALLBACK",
            "to_zone_id": central_zone,
            "positive_direction": "from_to_positive",
            "purpose": "computational_routing",
        }
    )
    accepted_by_id = {
        str(item["boundary_id"]): item
        for item in review_candidate["accepted_cutsets"]
    }
    boundaries: list[dict[str, object]] = []
    for proposal in sorted(cut_proposals, key=lambda item: str(item.get("boundary_id"))):
        boundary_id = str(proposal.get("boundary_id") or "")
        capability = capabilities.get(boundary_id)
        if capability is None:
            raise ValueError(f"Selected ETYS boundary {boundary_id} lacks 2025 capability evidence")
        accepted = accepted_by_id[boundary_id]
        boundaries.append(
            {
                "boundary_id": boundary_id,
                "display_name": f"ETYS {boundary_id}",
                "source_partition_zone_ids": list(accepted["positive_zone_ids"]),
                "sink_partition_zone_ids": list(accepted["negative_zone_ids"]),
                "forward_limit_mw": float(capability["forward_limit_mw"]),
                "source_direction": "north_or_positive_side_to_south_or_negative_side",
                "expected_corridor_ids": sorted(
                    str(item["corridor_id"]) for item in accepted["members"]
                ),
            }
        )
    return {
        "schema_version": "value.prompt98-provisional-network-payload/v1",
        "dso_geojson": {
            "type": "FeatureCollection",
            "features": zone_features + [fallback],
        },
        "etys_payload": {
            "corridors": sorted(corridors, key=lambda item: str(item["corridor_id"])),
            "boundaries": boundaries,
            "rating_profiles": [],
            "unconstrained_zone_ids": ["ENGLAND_FALLBACK"],
        },
        "review": {
            "owner_signoff": "pending",
            "provisional_mapping_approval": "builder_only",
            "fallback_transport_zone": central_zone,
            "corridor_semantics": "computational_routing_not_physical_circuits",
            "cut_orientation_method": "centroid_north_of_nearest_boundary_point",
            "candidate": review_candidate,
        },
    }


def solar_availability(values: Sequence[float]) -> list[float]:
    """Copied VALUE irradiance-to-output equation, expressed per MW."""

    result: list[float] = []
    for raw in values:
        value = float(raw)
        result.append(value / 3_600_000.0 if 3_600 < value <= 36_000_000 else 0.0)
    return result


def wind_availability(values: Sequence[float], *, technology: str) -> list[float]:
    """Copied VALUE cubic wind curve, normalized from its 20 MW unit."""

    if technology == "onshore":
        rated, cut_out = 9.7, 25.0
    elif technology == "offshore":
        rated, cut_out = 10.5, 30.0
    else:
        raise ValueError("Wind availability requires onshore or offshore technology")
    result: list[float] = []
    for raw in values:
        speed = float(raw)
        if speed < 3.0 or speed > cut_out:
            result.append(0.0)
        elif speed >= rated:
            result.append(1.0)
        else:
            result.append(max(0.0, min(1.0, (speed**3 - 3.0**3) / (rated**3 - 3.0**3))))
    return result


def bounded_availability(
    values: Sequence[float], *, tolerance: float = 1e-12
) -> list[float]:
    """Close floating residue only; reject a scientifically material overshoot."""

    result: list[float] = []
    for raw in values:
        value = float(raw)
        if not math.isfinite(value) or value < -tolerance or value > 1.0 + tolerance:
            raise ValueError("Compiled availability is materially outside 0..1")
        result.append(min(1.0, max(0.0, value)))
    return result


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _half_hour(values: Sequence[float], periods: int) -> list[float]:
    import numpy as np

    hourly = np.asarray(values, dtype=float).reshape(-1)
    doubled = np.repeat(hourly, 2)
    if len(doubled) < periods:
        doubled = np.tile(doubled, int(math.ceil(periods / max(len(doubled), 1))))
    result = doubled[:periods].tolist()
    if any(not math.isfinite(value) or not 0 <= value <= 1 for value in result):
        raise ValueError("Compiled ERA5 availability is outside 0..1")
    return result


def build_weather_bundles(
    fleet_rows: Sequence[Mapping[str, object]],
    repd_rows: Sequence[Mapping[str, object]],
    zone_geojson: Mapping[str, object],
    *,
    wind_nc: Path,
    solar_nc: Path,
    period_ids: Sequence[str],
) -> dict[str, object]:
    """Build both Prompt 97 weather modes entirely offline."""

    import numpy as np
    import xarray as xr

    periods = tuple(str(item) for item in period_ids)
    wind_hash = _file_sha256(Path(wind_nc))
    solar_hash = _file_sha256(Path(solar_nc))
    combined_hash = hashlib.sha256(f"{wind_hash}:{solar_hash}".encode("ascii")).hexdigest()
    wind_data = xr.open_dataset(Path(wind_nc))
    solar_data = xr.open_dataset(Path(solar_nc))
    try:
        cache: dict[tuple[str, float, float], list[float]] = {}

        def profile(technology: str, latitude: float, longitude: float) -> list[float]:
            dataset = solar_data if technology == "solar" else wind_data
            latitude_value = float(dataset["latitude"].sel(latitude=latitude, method="nearest"))
            longitude_value = float(dataset["longitude"].sel(longitude=longitude, method="nearest"))
            key = (technology, latitude_value, longitude_value)
            if key in cache:
                return cache[key]
            variable = "ssrd" if technology == "solar" else "wind_speed"
            data = dataset[variable].sel(
                latitude=latitude_value,
                longitude=longitude_value,
            )
            order = [name for name in ("dayofyear", "hour") if name in data.dims]
            if set(order) != {"dayofyear", "hour"}:
                raise ValueError(f"ERA5 {variable} lacks dayofyear/hour dimensions")
            values = data.transpose(*order).values.reshape(-1)
            hourly = (
                solar_availability(values.tolist())
                if technology == "solar"
                else wind_availability(values.tolist(), technology=technology)
            )
            cache[key] = _half_hour(hourly, len(periods))
            return cache[key]

        active_vre = [
            row for row in fleet_rows
            if str(row.get("status") or "").casefold() in {"operating", "operational", "active", "commissioned"}
            and str(row.get("technology") or "") in {"solar", "onshore", "offshore"}
            and float(row.get("capacity_mw") or 0.0) > 0
        ]
        representative_profiles: list[ZonalAvailabilityProfile] = []
        for row in sorted(active_vre, key=lambda item: str(item.get("asset_id"))):
            if row.get("latitude") is None or row.get("longitude") is None:
                raise ValueError(f"VRE asset {row.get('asset_id')} lacks its representative coordinate")
            technology = str(row["technology"])
            asset_id = str(row["asset_id"])
            values = profile(technology, float(row["latitude"]), float(row["longitude"]))
            representative_profiles.append(
                ZonalAvailabilityProfile(
                    f"representative:{asset_id}:{technology}:{row['zone_id']}",
                    asset_id,
                    technology,
                    str(row["zone_id"]),
                    periods,
                    values,
                    float(row["capacity_mw"]),
                    (asset_id,),
                    "copied_scheme_c_representative_coordinate_trace",
                    solar_hash if technology == "solar" else wind_hash,
                )
            )

        repd_groups: dict[tuple[str, str], list[Mapping[str, object]]] = defaultdict(list)
        all_by_technology: dict[str, list[Mapping[str, object]]] = defaultdict(list)
        for row in repd_rows:
            technology = str(row.get("technology") or "")
            if (
                str(row.get("status") or "").casefold() != "operational"
                or technology not in {"solar", "onshore", "offshore"}
                or row.get("latitude") is None
                or row.get("longitude") is None
            ):
                continue
            zone_id, _ = _nearest_zone(
                float(row["longitude"]), float(row["latitude"]), zone_geojson
            )
            if not zone_id:
                continue
            enriched = {**dict(row), "compiled_zone_id": zone_id}
            repd_groups[(technology, zone_id)].append(enriched)
            all_by_technology[technology].append(enriched)

        aggregated_profiles: list[ZonalAvailabilityProfile] = []
        fallback_groups: list[dict[str, object]] = []
        for row in sorted(active_vre, key=lambda item: str(item.get("asset_id"))):
            technology = str(row["technology"])
            zone_id = str(row["zone_id"])
            sources = repd_groups.get((technology, zone_id), [])
            method = "offline_repd_location_era5_mw_weighted"
            if not sources:
                sources = all_by_technology.get(technology, [])
                method = "gb_technology_repd_fallback_no_same_zone_sources"
                fallback_groups.append(
                    {
                        "asset_id": str(row["asset_id"]),
                        "technology": technology,
                        "zone_id": zone_id,
                        "reason": "no_operational_repd_sources_in_zone",
                    }
                )
            if sources:
                total = sum(float(item["capacity_mw"]) for item in sources)
                weighted = np.zeros(len(periods), dtype=float)
                for source in sources:
                    values = profile(
                        technology,
                        float(source["latitude"]),
                        float(source["longitude"]),
                    )
                    weighted += np.asarray(values) * float(source["capacity_mw"]) / total
                source_ids = tuple(sorted(str(item["asset_id"]) for item in sources))
                values = bounded_availability(weighted.tolist())
            else:
                values = profile(technology, float(row["latitude"]), float(row["longitude"]))
                source_ids = (str(row["asset_id"]),)
                method = "representative_fallback_no_operational_repd_technology"
            aggregated_profiles.append(
                ZonalAvailabilityProfile(
                    f"repd-era5:{row['asset_id']}:{technology}:{zone_id}",
                    str(row["asset_id"]),
                    technology,
                    zone_id,
                    periods,
                    values,
                    float(row["capacity_mw"]),
                    source_ids,
                    method,
                    solar_hash if technology == "solar" else wind_hash,
                )
            )
    finally:
        wind_data.close()
        solar_data.close()

    representative = ZonalAvailabilityBundle(
        "representative_point",
        periods,
        tuple(representative_profiles),
        combined_hash,
        runtime_opens_repd_or_era5=False,
        provenance={
            "equations": "copied_scheme_c_solar_and_wind_availability",
            "period_conversion": "hourly_values_repeated_twice_then_truncated_to_accepted_clock",
        },
    )
    representative = replace(
        representative, scientific_sha256=representative.compute_scientific_sha256()
    )
    representative.validate()
    aggregated = ZonalAvailabilityBundle(
        "repd_era5_mw_aggregated",
        periods,
        tuple(aggregated_profiles),
        combined_hash,
        runtime_opens_repd_or_era5=False,
        provenance={
            "aggregation": "operational_repd_mw_weighted_by_technology_and_zone",
            "fallback_groups": fallback_groups,
            "runtime_inputs": "aggregated_profiles_only",
        },
    )
    aggregated = replace(
        aggregated, scientific_sha256=aggregated.compute_scientific_sha256()
    )
    aggregated.validate()
    return {
        "schema_version": "value.prompt98-weather-build/v1",
        "representative": representative.to_dict(),
        "aggregated": aggregated.to_dict(),
        "fallback_groups": fallback_groups,
        "source_sha256": {"wind": wind_hash, "solar": solar_hash},
    }


_INTERCONNECTOR_COUNTRY = {
    "auchencrosh (interconnector cct) *": "Ireland",
    "britned": "Netherlands",
    "east west interconnector": "Ireland",
    "eleclink": "France",
    "greenlink": "Ireland",
    "ifa interconnector": "France",
    "ifa2 interconnector": "France",
    "nemo link": "Belgium",
    "ns link": "Norway",
}

_INTERCONNECTOR_SITE_COORDINATES = {
    "Auchencrosh 275kV": (-4.95, 55.16),
    "Grain 400kV Substation": (0.71, 51.45),
    "Deeside 400kV Substation": (-3.06, 53.23),
    "Sellindge 400kV Substation": (1.00, 51.08),
    "Pembroke 400kV Substation": (-4.99, 51.68),
    "Chilling 400kV Substation": (-1.27, 50.82),
    "Richborough 400kV Substation": (1.34, 51.31),
    "Blyth GSP": (-1.51, 55.13),
}


def normalize_interconnector_register(
    source: Path,
    zone_geojson: Mapping[str, object],
) -> dict[str, object]:
    """Compile built links used by the five retained VALUE country profiles."""

    by_project: dict[str, list[dict[str, str]]] = defaultdict(list)
    excluded: dict[str, int] = defaultdict(int)
    with Path(source).open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            project = str(row.get("Project Name") or "").strip()
            if str(row.get("Project Status") or "").strip().casefold() != "built":
                excluded["not_built"] += 1
                continue
            if project.casefold() not in _INTERCONNECTOR_COUNTRY:
                excluded["outside_scheme_c_country_profile_scope"] += 1
                continue
            by_project[project].append(dict(row))
    rows: list[dict[str, object]] = []
    for project in sorted(by_project, key=str.casefold):
        candidates = by_project[project]

        def stage_key(row: Mapping[str, object]) -> tuple[float, str]:
            try:
                stage = float(row.get("Stage") or 0.0)
            except (TypeError, ValueError):
                stage = 0.0
            return stage, str(row.get("MW Effective From") or "")

        selected = max(candidates, key=stage_key)
        site = str(selected.get("Connection Site") or "").strip()
        coordinate = _INTERCONNECTOR_SITE_COORDINATES.get(site)
        if coordinate is None:
            raise ValueError(
                f"Built VALUE interconnector {project} lacks reviewed landing-site evidence"
            )
        zone_id, covered = _nearest_zone(coordinate[0], coordinate[1], zone_geojson)
        if not zone_id:
            raise ValueError(f"Built VALUE interconnector {project} has no GB landing zone")
        try:
            import_capacity = float(selected.get("MW Import - Current") or 0.0)
            export_capacity = float(selected.get("MW Export - Current") or 0.0)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Built VALUE interconnector {project} has invalid capability") from exc
        if import_capacity <= 0 and export_capacity <= 0:
            raise ValueError(f"Built VALUE interconnector {project} has zero current capability")
        asset_id = f"interconnector:{_slug(project)}"
        rows.append(
            {
                "interconnector_id": asset_id,
                "asset_id": asset_id,
                "project_name": project,
                "economic_country_agent": f"Interconnect_{_INTERCONNECTOR_COUNTRY[project.casefold()]}",
                "counterparty_country": _INTERCONNECTOR_COUNTRY[project.casefold()],
                "connection_site": site,
                "zone_id": zone_id,
                "longitude": coordinate[0],
                "latitude": coordinate[1],
                "landing_method": (
                    "curated_connection_site_coordinate"
                    if covered else "nearest_coast_dso_inference"
                ),
                "landing_evidence_level": "candidate_crosswalk_pending_owner_review",
                "import_capacity_mw": import_capacity,
                "export_capacity_mw": export_capacity,
                "capacity_mw": max(import_capacity, export_capacity),
                "technology": "interconnector",
                "status": "operating",
                "envelope_semantics": "signed_profile_envelope",
                "effective_date": str(selected.get("MW Effective From") or ""),
                "selected_stage": str(selected.get("Stage") or ""),
                "source_stage_count": len(candidates),
            }
        )
    return {
        "schema_version": "value.prompt98-interconnector-landings/v1",
        "rows": rows,
        "excluded_by_reason": dict(sorted(excluded.items())),
        "country_profile_semantics": "physical_links_are_tranches_of_five_retained_country_agents",
        "owner_signoff": "pending",
    }
