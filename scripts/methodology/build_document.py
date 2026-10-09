from pathlib import Path
import json,re,base64,sys,os
from lxml import html,etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
W=Path(os.environ.get('VALUE_METHODOLOGY_WORK',Path(__file__).resolve().parent));O=Path(os.environ.get('VALUE_DOCUMENT_DIR',str(W/'output')));O.mkdir(parents=True,exist_ok=True)
lang=sys.argv[1] if len(sys.argv)>1 else 'zh'
suffix='-en' if lang=='en' else ''
mathdir='math'+suffix
source=Path(os.environ.get('VALUE_METHODOLOGY_SOURCE',W))
edition=json.loads((source/'edition.json').read_text())
stem=('VALUE_Model_Methodology_EN_' if lang=='en' else 'VALUE_模型方法学_')+edition['date']
offline=('VALUE_methodology_EN_' if lang=='en' else 'VALUE_methodology_ZH_')+edition['date']+'.html'
title='VALUE Model Methodology' if lang=='en' else 'VALUE 模型方法学'
subtitle='Mathematical formulation  Algorithms  Input data' if lang=='en' else '数学设定  计算方法  输入数据'
revision=edition['revision'][lang]
basis=edition['basisLabel'][lang]
payload=json.loads((W/('rendered'+suffix+'.json')).read_text());eq={x['id']:x for x in payload['equations']}
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.67);sec.right_margin=Inches(.67)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)
styles=doc.styles
for border in list(styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].element.get_or_add_rPr().append(OxmlElement('w:lang'))
doc.styles['Normal'].element.rPr[-1].set(qn('w:val'),'en-GB')
doc.styles['Normal'].element.rPr[-1].set(qn('w:eastAsia'),'zh-CN')
for st in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Heading 4','Table Grid']:
 s=styles[st];s.font.name='Liberation Serif';s.font.color.rgb=RGBColor(0,0,0);s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
styles['Normal'].font.size=Pt(11.5 if lang=='en' else 11)
styles['Normal'].paragraph_format.line_spacing=1.23;styles['Normal'].paragraph_format.space_after=Pt(6)
if lang=='zh':
 for name in ['Normal','Table Grid','Heading 1','Heading 2','Heading 3','Heading 4']:
  pr=styles[name].element.get_or_add_pPr()
  for tag,value in [('wordWrap','0'),('kinsoku','1')]:
   flag=OxmlElement('w:'+tag);flag.set(qn('w:val'),value);pr.append(flag)
for name,size in [('Title',25),('Subtitle',12),('Heading 1',19),('Heading 2',14),('Heading 3',11.5),('Heading 4',10.5)]:
 styles[name].font.size=Pt(size);styles[name].paragraph_format.space_before=Pt(13);styles[name].paragraph_format.space_after=Pt(7)
 styles[name].paragraph_format.keep_with_next=True
styles['Title'].paragraph_format.space_before=Pt(0)
styles['Heading 1'].paragraph_format.space_before=Pt(24)
sec.header.paragraphs[0].text=title
sec.header.paragraphs[0].style='Normal'
for r in sec.header.paragraphs[0].runs:r.font.size=Pt(8)
f=sec.footer.paragraphs[0];f.alignment=2;f.add_run(revision+'    ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
for r in f.runs:r.font.size=Pt(8)
doc.core_properties.title=title;doc.core_properties.subject=subtitle;doc.core_properties.author='VALUE'
doc.add_paragraph(title,'Title')
doc.add_paragraph(subtitle,'Subtitle')
doc.add_paragraph(revision+'\n'+basis)
doc.add_paragraph('Contents' if lang=='en' else '目录','Heading 2')
for ch in payload['chapters']:
    heading=re.search(r'<h2>(.*?)</h2>',ch['html']).group(1)
    doc.add_paragraph(re.sub('<[^>]+>','',heading))
def soft(s):
 # Break long identifiers while keeping numerical units intact.
 return re.sub(r'[^\s]{24,}',lambda m:m.group().replace('_','_\u200b').replace('/','/\u200b'),s)
def addtext(p,text,bold=False,italic=False,code=False,small=False):
 if not text:return
 r=p.add_run(soft(text));r.bold=bold;r.italic=italic
 if code:r.font.name='DejaVu Sans Mono';r.font.size=Pt(8.5)
 if small:r.font.size=Pt(9)
 return r
def inline(p,node,bold=False,italic=False,code=False,small=False):
 if node.text:addtext(p,node.text,bold,italic,code,small)
 for c in node:
  if c.tag=='img' and c.get('data-equation'):
   e=eq[c.get('data-equation')];w=e['width']/96*(11.5/12 if lang=='en' else 11/12);p.add_run().add_picture(str(W/mathdir/f"{e['id']}.png"),width=Inches(min(w,5.8)))
  elif c.tag=='br':p.add_run().add_break()
  else:inline(p,c,bold or c.tag in ['b','strong'],italic or c.tag=='em',code or c.tag=='code',small)
  if c.tail:
   tail=c.tail
   if lang=='zh' and c.tag=='img' and re.match(r'^[，。；：、）]',tail):tail='\u2060'+tail
   addtext(p,tail,bold,italic,code,small)
def render_node(node,parent=doc):
 tag=node.tag
 if tag in ['h2','h3','h4','h5','h6']:
  p=parent.add_paragraph(style='Heading '+str(min(int(tag[1])-1,4)));inline(p,node)
 elif tag=='p':inline(parent.add_paragraph(),node)
 elif tag=='div' and node.get('data-equation'):
  e=eq[node.get('data-equation')];p=parent.add_paragraph();p.alignment=1;p.paragraph_format.space_after=Pt(8)
  p.add_run().add_picture(str(W/mathdir/f"{e['id']}.png"),width=Inches(min(e['width']/96,6.8)))
 elif tag=='pre':
  p=parent.add_paragraph();p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(8);p.paragraph_format.keep_together=True
  inline(p,node,code=True)
 elif tag in ['ul','ol']:
  for i,li in enumerate(node.findall('li'),1):
   p=parent.add_paragraph();p.paragraph_format.left_indent=Inches(.12)
   p.add_run(f'{i}. ' if tag=='ol' else '• ')
   inline(p,li)
 elif tag=='table':
  rows=node.xpath('./thead/tr|./tbody/tr|./tr');n=max(len(r) for r in rows)
  table=parent.add_table(rows=0,cols=n);table.style='Table Grid';table.autofit=False
  total=6.8
  weights=[]
  for ci in range(n):
   lengths=[sum(1.8 if ord(ch)>255 else 1 for ch in r[ci].text_content()) for r in rows if len(r)>ci]
   weights.append(max(8,min(38,max(lengths,default=8)**.7*2.5)))
  widths=[.62+(total-.62*n)*v/sum(weights) for v in weights]
  for i,col in enumerate(table.columns):col.width=Inches(widths[i])
  for ri,row in enumerate(rows):
   cells=table.add_row().cells
   trpr=table.rows[-1]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
   if ri==0:
    repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trpr.append(repeat)
   for ci,source in enumerate(row):
    cell=cells[ci];cell.width=Inches(widths[ci]);cell.vertical_alignment=1
    pr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
    for side in ['top','left','bottom','right']:
     x=OxmlElement('w:'+side);x.set(qn('w:w'),'60' if ch['id']=='appendix' and side in ['top','bottom'] else '90');x.set(qn('w:type'),'dxa');margins.append(x)
    pr.append(margins)
    if ri==0:
     sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');pr.append(sh)
    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.15
    inline(p,source,bold=ri==0,small=True)
   # Explicit light neutral borders.
  pr=table._tbl.tblPr;b=OxmlElement('w:tblBorders')
  for side in ['top','left','bottom','right','insideH','insideV']:
   x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
  pr.append(b);parent.add_paragraph().paragraph_format.space_after=Pt(1)
 elif tag not in ['hr']:
  if node.text and node.text.strip():addtext(parent.add_paragraph(),node.text)
  for c in node:render_node(c,parent)
for ch_index,ch in enumerate(payload['chapters']):
 if ch_index==0:doc.add_page_break()
 root=html.fragment_fromstring(ch['html'],create_parent='div')
 for node_index,node in enumerate(root):
  render_node(node)
  if node.tag=='p' and node_index+1<len(root) and root[node_index+1].get('data-equation'):
   doc.paragraphs[-1].paragraph_format.keep_with_next=True
for border in list(doc.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.save(O/(stem+'.docx'))
toc=''.join(f'<a href="#{ch["id"]}">{re.search(r"<h2>(.*?)</h2>",ch["html"])[1]}</a>' for ch in payload['chapters'])
body=''.join(f'<section id="{ch["id"]}">{ch["html"]}</section>' for ch in payload['chapters'])
for e in payload['equations']:
 data=base64.b64encode((W/mathdir/f"{e['id']}.svg").read_bytes()).decode()
 body=body.replace(mathdir+'/'+e['id']+'.svg','data:image/svg+xml;base64,'+data)
css='''body{font-family:Georgia,'Noto Serif CJK SC',serif;max-width:960px;margin:48px auto;padding:0 24px;color:#182235;line-height:1.8}h1{font-size:32px}h2{margin-top:64px;font-size:27px;border-bottom:1px solid #ccc;padding-bottom:16px}h3{margin-top:40px}table{border-collapse:collapse;width:100%;font-size:14px;margin:24px 0}td,th{border:1px solid #cbd0d4;padding:10px;text-align:left;overflow-wrap:anywhere}th{background:#eef2f2}code,pre{font-family:"DejaVu Sans Mono",monospace;font-size:.88em;overflow-wrap:anywhere}pre{white-space:pre-wrap;background:#f4f6f6;padding:18px}nav a{display:block;padding:4px 0;color:#236c69}.equation{max-width:100%;overflow-x:auto;margin:24px 0;text-align:center}.equation img{max-width:100%;height:auto}.inline-equation{height:1.4em;vertical-align:middle;max-width:100%}a{color:#146663}section{scroll-margin-top:30px}@media(max-width:600px){body{padding:0 16px}table{display:block;overflow-x:auto}h1{font-size:26px}}'''
(O/offline).write_text('<!doctype html><html lang="'+lang+'"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+title+'</title><style>'+css+'</style></head><body><h1>'+title+'</h1><p>'+revision+' · '+basis+'</p><nav aria-label="Contents">'+toc+'</nav>'+body+'</body></html>')
print('DOCX and standalone HTML created',len(payload['chapters']),'chapters',len(eq),'equations')
